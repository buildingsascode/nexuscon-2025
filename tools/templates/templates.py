from pathlib import Path

import openpyxl

from tools.templates.points import PointTemplate

TEMPLATE_PATH = Path("tools/templates/point_template.xlsx")


def get_template_points() -> list[PointTemplate]:
    """Load the point template and return a list of PointTemplate instances."""
    with open(TEMPLATE_PATH, "rb") as f:
        template_workbook = openpyxl.load_workbook(f)
        template_worksheet = template_workbook["Switch"]
        rows = []
        for row in template_worksheet.iter_rows(
            min_row=2, max_row=template_worksheet.max_row, values_only=True
        ):
            row_dict = {
                template_worksheet.cell(row=1, column=i + 1).value: value
                for i, value in enumerate(row)
            }
            rows.append(PointTemplate.model_validate(row_dict))
        return rows
